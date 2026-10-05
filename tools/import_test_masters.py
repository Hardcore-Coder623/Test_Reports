"""Import the model nameplate master and saved test readings from the Excel test-report workbook.

Usage:  python tools/import_test_masters.py "Test_Report_Program_Tech_Auto_final.xlsm"
Writes data/nameplates.json (sheet NameplateData) and data/test_records.json (sheet TestData),
replacing what is there. Records saved from the web page are lost unless they are in the workbook.
"""
import datetime
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
KEYMAP = {"Model": "model", "No_of_Stages": "stages", "Delivery_Size": "delivery_size", "Head(HG)": "head_g",
          "Discharge(QG)": "q_g", "Discharge_Unit": "q_unit", "OverallEff%": "eff_g",
          "Min_Submergence": "min_sub", "Type_of_Duty": "duty_type", "Motor_Rating_kW": "motor_kw",
          "Motor_Rating_HP": "motor_hp", "Rated_Speed": "rated_speed", "Rated_Current": "rated_current",
          "Rated_Voltage": "rated_voltage", "Category": "category", "Rated_Freq": "rated_freq",
          "Winding_Conn": "winding", "Type_of_Motor": "motor_type", "No_of_phase": "phases",
          "Operating_H-Min": "h_min", "Operating-H-Max": "h_max", "Brand": "brand", "ISI": "isi"}
READING_FIELDS = ["z", "h", "q", "v", "r", "y", "b", "w", "hz", "n"]


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["NameplateData"]
    hdr = [c.value for c in ws[1]]
    nps = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        d = {KEYMAP[h]: ("" if v is None else v) for h, v in zip(hdr, row) if h in KEYMAP}
        d["model"] = str(d["model"]).strip()
        nps.append(d)

    ws = wb["TestData"]
    recs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        date = row[1]
        if isinstance(date, (datetime.datetime, datetime.date)):
            date = date.strftime("%Y-%m-%d")
        readings = []
        for k in range(10):                                   # Reading-1 .. Reading-10, 10 columns each
            vals = row[3 + k * 10: 3 + k * 10 + 10]
            if len(vals) < 3 or (vals[1] is None and vals[2] is None):
                continue
            readings.append({f: ("" if v is None else v) for f, v in zip(READING_FIELDS, vals)})
        if not isinstance(date, str) or date.startswith("#"):        # blank or Excel error such as #N/A
            date = ""
        recs.append({"sr_no": str(row[0]).strip(), "date": date or "", "model": str(row[2] or "").strip(),
                     "readings": readings})

    (ROOT / "data" / "nameplates.json").write_text(json.dumps(nps, indent=0, default=str), encoding="utf-8")
    (ROOT / "data" / "test_records.json").write_text(json.dumps(recs, separators=(",", ":"), default=str),
                                                     encoding="utf-8")
    print(f"Imported {len(nps)} nameplates and {len(recs)} test records.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
